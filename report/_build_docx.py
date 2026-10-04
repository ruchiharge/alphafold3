"""Convert report/FinalReport.md into report/Final_Report_OptimalAntibodyAntigen.docx.

The output is formatted to match the institutional template used in the
mid-semester report (Midsem_Report_OptimalAntibodyAntigen.docx):

- Centered title page with the BITS emblem at the top
- Acknowledgments page
- Certificate page
- Abstract + Contents + all numbered sections (taken from the markdown)
- Times New Roman throughout, centered Heading 1 section titles

Markdown features handled (a strict subset, matching what FinalReport.md uses):
- # / ## / ### headings
- bold (**...**), italic (*...*) and inline code (`...`)
- bullet lists (- ...)
- numbered lists (1. ...)
- code fences (```lang ... ```)
- pipe tables (| col | col |)
- images (![alt](path))
- raw HTML videos (<video controls src="x.mp4" ...></video>) — replaced with x_frame.png
- horizontal rules (---) become a blank paragraph
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parent
SOURCE_MD = REPORT_DIR / "FinalReport.md"
OUTPUT_DOCX = REPORT_DIR / "Final_Report_OptimalAntibodyAntigen.docx"
BITS_EMBLEM = REPORT_DIR / "BITSemblem.jpg"

DEFAULT_FONT = "Times New Roman"
DEFAULT_SIZE = Pt(12)

INLINE_CODE_RE = re.compile(r"`([^`]+)`")
BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")
ITALIC_RE = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
VIDEO_RE = re.compile(r'<video[^>]*src="([^"]+)"[^>]*></video>', re.IGNORECASE)
TABLE_ROW_RE = re.compile(r"^\s*\|.*\|\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?(?:\s*\|\s*:?-+:?)+\s*\|?\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
ULIST_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
OLIST_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
HR_RE = re.compile(r"^\s*---+\s*$")
HTML_COMMENT_OPEN = "<!--"
HTML_COMMENT_CLOSE = "-->"


def _set_run_default_font(run, *, size: Pt = DEFAULT_SIZE, bold: bool | None = None,
                           italic: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = DEFAULT_FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), DEFAULT_FONT)
    rFonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    rFonts.set(qn("w:cs"), DEFAULT_FONT)
    rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_runs(paragraph, text: str, *, base_size: Pt = DEFAULT_SIZE,
             base_bold: bool | None = None, base_italic: bool | None = None) -> None:
    """Render inline markdown (**bold**, *italic*, `code`, [text](url), ![alt](path)) into runs."""
    pos = 0
    pattern = re.compile(
        r"(\*\*[^*]+\*\*)"             # bold
        r"|(\*[^*]+\*)"                # italic
        r"|(`[^`]+`)"                  # code
        r"|(!\[[^\]]*\]\([^)]+\))"     # image (emit alt text only here)
        r"|(\[[^\]]+\]\([^)]+\))"      # link (emit label only)
    )
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            _set_run_default_font(run, size=base_size, bold=base_bold, italic=base_italic)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_default_font(run, size=base_size, bold=True, italic=base_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = base_size
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            _set_run_default_font(run, size=base_size, bold=base_bold, italic=True)
        elif token.startswith("!["):
            alt = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", token).group(1)
            run = paragraph.add_run(alt)
            _set_run_default_font(run, size=base_size, bold=base_bold, italic=base_italic)
        elif token.startswith("["):
            label = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).group(1)
            run = paragraph.add_run(label)
            _set_run_default_font(run, size=base_size, bold=base_bold, italic=base_italic)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_default_font(run, size=base_size, bold=base_bold, italic=base_italic)


def add_image(doc: Document, image_path: Path, *, max_width_in: float = 6.0) -> None:
    if not image_path.exists():
        p = doc.add_paragraph(f"[missing image: {image_path.name}]")
        for r in p.runs:
            _set_run_default_font(r)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    try:
        run.add_picture(str(image_path), width=Inches(max_width_in))
    except Exception as exc:
        err = doc.add_paragraph(f"[failed to insert {image_path.name}: {exc}]")
        for r in err.runs:
            _set_run_default_font(r)


def add_code_block(doc: Document, lang: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F4F4")
    pPr.append(shd)
    text = "\n".join(lines)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        cell = hdr_cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h.strip())
        _set_run_default_font(run, size=Pt(11), bold=True)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, val.strip(), base_size=Pt(11))


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def video_to_frame_path(src: str) -> Path:
    p = Path(src)
    return REPORT_DIR / (p.stem + "_frame.png")


def resolve_image_path(src: str) -> Path:
    p = Path(src)
    if p.is_absolute():
        return p
    return REPORT_DIR / p


# ----- Document set-up & front matter ----------------------------------------------

def configure_default_styles(doc: Document) -> None:
    """Force Times New Roman 12 pt as the default body font and sane heading sizes."""
    normal = doc.styles["Normal"]
    normal.font.name = DEFAULT_FONT
    normal.font.size = DEFAULT_SIZE
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), DEFAULT_FONT)
    rFonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    rFonts.set(qn("w:cs"), DEFAULT_FONT)
    rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)

    heading_sizes = {"Heading 1": Pt(16), "Heading 2": Pt(13), "Heading 3": Pt(12), "Heading 4": Pt(11)}
    for name, size in heading_sizes.items():
        try:
            s = doc.styles[name]
            s.font.name = DEFAULT_FONT
            s.font.size = size
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
        except KeyError:
            pass


def add_centered_paragraph(doc: Document, text: str, *, size: Pt = DEFAULT_SIZE,
                            bold: bool | None = None, italic: bool | None = None,
                            space_after: Pt | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    _set_run_default_font(run, size=size, bold=bold, italic=italic)


def add_blank(doc: Document, n: int = 1) -> None:
    for _ in range(n):
        doc.add_paragraph()


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_title_page(doc: Document) -> None:
    if BITS_EMBLEM.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        try:
            run.add_picture(str(BITS_EMBLEM), width=Inches(1.6))
        except Exception:
            pass

    add_centered_paragraph(doc, "A Project Report", size=Pt(14), italic=True)
    add_centered_paragraph(doc, "On", size=Pt(12))
    add_centered_paragraph(
        doc,
        "Modeling Antibody–Antigen Interactions Using AlphaFold3",
        size=Pt(16),
        bold=True,
    )
    add_blank(doc)
    add_centered_paragraph(doc, "BY", size=Pt(12))
    add_centered_paragraph(doc, "SHUBH SRIVASTAVA   2022B4A70664H", size=Pt(13), bold=True)
    add_centered_paragraph(doc, "ATHARVA CHAPHEKAR   2022B4A71815H", size=Pt(13), bold=True)
    add_centered_paragraph(doc, "RUCHI HARGE   2022B4A70942H", size=Pt(13), bold=True)
    add_centered_paragraph(doc, "RITHIKA BALUSU   2022B1A70938H", size=Pt(13), bold=True)
    add_blank(doc)
    add_centered_paragraph(doc, "Under the supervision of", size=Pt(12))
    add_centered_paragraph(
        doc,
        "Prof. Prajakta Bedekar and Prof. Satyanarayana Subrahmanyam",
        size=Pt(13),
        bold=True,
    )
    add_blank(doc, 2)
    add_centered_paragraph(
        doc,
        "SUBMITTED IN PARTIAL FULFILLMENT OF THE REQUIREMENTS OF",
        size=Pt(12),
        bold=True,
    )
    add_centered_paragraph(doc, "MATH F376: DESIGN PROJECT", size=Pt(12), bold=True)
    add_blank(doc)
    add_centered_paragraph(
        doc,
        "BIRLA INSTITUTE OF TECHNOLOGY AND SCIENCE PILANI (RAJASTHAN)",
        size=Pt(12),
        bold=True,
    )
    add_centered_paragraph(doc, "HYDERABAD CAMPUS", size=Pt(12), bold=True)
    add_centered_paragraph(doc, "(APRIL 2026)", size=Pt(12), bold=True)

    add_page_break(doc)


def add_acknowledgments_page(doc: Document) -> None:
    add_centered_paragraph(doc, "ACKNOWLEDGMENTS", size=Pt(16), bold=True)
    add_blank(doc)
    body = (
        "We sincerely express our profound gratitude to "
        "Prof. Prajakta Bedekar and Prof. Satyanarayana Subrahmanyam for providing us "
        "with the opportunity to undertake this Design Project and for their continued "
        "guidance throughout its course. Their advice on framing the scientific question, "
        "on selecting an antibody–antigen system worth studying, and on choosing a "
        "saturation–mutagenesis strategy that fit within our HPC budget shaped this work "
        "decisively."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, body)

    body2 = (
        "We are equally grateful to the BITS Hyderabad High Performance Computing facility "
        "for granting access to the GPU partitions on which every AlphaFold3 prediction in this "
        "report was executed, and to the support staff who maintain the SLURM scheduler and the "
        "Singularity image we relied on. Without their infrastructure the saturation scans at "
        "VEGFA Glu 74 (demo6 and the ongoing demo7) would not have been feasible."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, body2)

    body3 = (
        "Finally, we thank the developers of AlphaFold3, ChimeraX and OpenMM for releasing the "
        "tools that make a project of this kind possible at the undergraduate level, and our "
        "peers for the many discussions that improved the analysis presented here."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, body3)

    add_page_break(doc)


def add_certificate_page(doc: Document) -> None:
    add_centered_paragraph(
        doc, "Birla Institute of Technology and Science-Pilani,", size=Pt(13), bold=True
    )
    add_centered_paragraph(doc, "Hyderabad Campus", size=Pt(13), bold=True)
    add_blank(doc)
    add_centered_paragraph(doc, "Certificate", size=Pt(16), bold=True)
    add_blank(doc)
    body = (
        "This is to certify that the project report entitled "
        "“Modeling Antibody–Antigen Interactions Using AlphaFold3” submitted by "
        "Shubh Srivastava (2022B4A70664H), Atharva Chaphekar (2022B4A71815H), "
        "Ruchi Harge (2022B4A70942H) and Rithika Balusu (2022B1A70938H) "
        "in partial fulfilment of the requirements of MATH F376: Design Project "
        "embodies the work done by them under our supervision and guidance."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, body)
    add_blank(doc, 4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Date: 28/04/26")
    _set_run_default_font(run, bold=True)
    run = p.add_run("\t\t\t\t\t\t\t")
    _set_run_default_font(run)
    run = p.add_run("(Prof. Prajakta Bedekar)")
    _set_run_default_font(run, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("BITS-Pilani, Hyderabad Campus")
    _set_run_default_font(run, bold=True)
    run = p.add_run("\t\t\t\t\t")
    _set_run_default_font(run)
    run = p.add_run("(Prof. Satyanarayana Subrahmanyam)")
    _set_run_default_font(run, bold=True)

    add_page_break(doc)


# ----- Markdown rendering ----------------------------------------------------------

def render_markdown_body(doc: Document, md_text: str) -> None:
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lang = ""
    code_buf: list[str] = []

    in_html_comment = False

    table_buf: list[list[str]] = []
    table_header: list[str] | None = None
    in_table = False

    def flush_table():
        nonlocal table_buf, table_header, in_table
        if table_header is not None and table_buf:
            add_table(doc, table_header, table_buf)
        table_buf = []
        table_header = None
        in_table = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip("\n")

        # HTML comments — skip whole block
        if HTML_COMMENT_OPEN in line and HTML_COMMENT_CLOSE in line:
            i += 1
            continue
        if HTML_COMMENT_OPEN in line:
            in_html_comment = True
            i += 1
            continue
        if in_html_comment:
            if HTML_COMMENT_CLOSE in line:
                in_html_comment = False
            i += 1
            continue

        # Fenced code blocks
        if line.startswith("```"):
            if not in_code:
                flush_table()
                in_code = True
                code_lang = line[3:].strip()
                code_buf = []
            else:
                add_code_block(doc, code_lang, code_buf)
                in_code = False
                code_lang = ""
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Tables — header row + separator + body rows
        if TABLE_ROW_RE.match(line):
            if i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].rstrip("\n")):
                flush_table()
                table_header = split_table_row(line)
                in_table = True
                i += 2
                while i < len(lines) and TABLE_ROW_RE.match(lines[i].rstrip("\n")):
                    table_buf.append(split_table_row(lines[i].rstrip("\n")))
                    i += 1
                flush_table()
                continue

        # Standalone image
        m_img_only = re.match(r"^\s*!\[[^\]]*\]\(([^)]+)\)\s*$", line)
        if m_img_only:
            flush_table()
            add_image(doc, resolve_image_path(m_img_only.group(1)))
            i += 1
            continue

        # Standalone <video> -> frame PNG
        m_vid = VIDEO_RE.search(line)
        if m_vid and re.match(r"^\s*<video", line):
            flush_table()
            frame = video_to_frame_path(m_vid.group(1))
            add_image(doc, frame)
            i += 1
            continue

        # Headings — promote one level up so that markdown ## (numbered sections in
        # FinalReport.md) becomes docx Heading 1, matching the midsem template where
        # "1. Introduction" is a top-level centered heading.
        m_h = HEADING_RE.match(line)
        if m_h:
            flush_table()
            level = len(m_h.group(1))
            text = m_h.group(2).strip()
            heading_level = max(level - 1, 1)
            heading_level = min(heading_level, 4)
            new_p = doc.add_heading("", level=heading_level)
            if heading_level == 1:
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            base_size = {1: Pt(16), 2: Pt(13), 3: Pt(12), 4: Pt(11)}.get(heading_level, Pt(12))
            add_runs(new_p, text, base_size=base_size, base_bold=True)
            i += 1
            continue

        # Horizontal rule
        if HR_RE.match(line):
            flush_table()
            doc.add_paragraph("")
            i += 1
            continue

        # Bullet list
        m_ul = ULIST_RE.match(line)
        if m_ul:
            flush_table()
            indent = len(m_ul.group(1))
            level = indent // 2
            p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            add_runs(p, m_ul.group(2))
            i += 1
            continue

        # Numbered list
        m_ol = OLIST_RE.match(line)
        if m_ol:
            flush_table()
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m_ol.group(3))
            i += 1
            continue

        # Blank line
        if line.strip() == "":
            flush_table()
            i += 1
            continue

        # Plain paragraph
        flush_table()
        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip("\n")
            if (
                nxt.strip() == ""
                or HEADING_RE.match(nxt)
                or HR_RE.match(nxt)
                or ULIST_RE.match(nxt)
                or OLIST_RE.match(nxt)
                or nxt.startswith("```")
                or TABLE_ROW_RE.match(nxt)
                or re.match(r"^\s*!\[[^\]]*\]\([^)]+\)\s*$", nxt)
                or re.match(r"^\s*<video", nxt)
                or HTML_COMMENT_OPEN in nxt
            ):
                break
            para_lines.append(nxt)
            i += 1
        joined = " ".join(s.strip() for s in para_lines)
        if IMAGE_RE.search(joined):
            parts = []
            pos = 0
            for m in IMAGE_RE.finditer(joined):
                if m.start() > pos:
                    parts.append(("text", joined[pos : m.start()]))
                parts.append(("image", m.group(2)))
                pos = m.end()
            if pos < len(joined):
                parts.append(("text", joined[pos:]))
            for kind, val in parts:
                if kind == "text":
                    if val.strip():
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        add_runs(p, val.strip())
                else:
                    add_image(doc, resolve_image_path(val))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, joined)

    flush_table()


def strip_markdown_front_matter(md_text: str) -> str:
    """Drop the title block at the top of FinalReport.md (everything up to and including the first `---`)."""
    lines = md_text.splitlines()
    for idx, ln in enumerate(lines):
        if HR_RE.match(ln):
            return "\n".join(lines[idx + 1 :])
    return md_text


def main() -> None:
    md = SOURCE_MD.read_text(encoding="utf-8")
    md_body = strip_markdown_front_matter(md)

    doc = Document()
    configure_default_styles(doc)

    add_title_page(doc)
    add_acknowledgments_page(doc)
    add_certificate_page(doc)
    render_markdown_body(doc, md_body)

    doc.save(str(OUTPUT_DOCX))
    print(f"Wrote {OUTPUT_DOCX} ({OUTPUT_DOCX.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
